import docx
from docx.enum.text import WD_BREAK
import os

def main():
    doc = docx.Document()
    p = doc.add_paragraph("First paragraph")
    p_ref = doc.add_paragraph("Second paragraph")
    
    # Let's insert a paragraph and a page break before p_ref
    p_new = p_ref.insert_paragraph_before("Inserted before second")
    p_break = p_ref.insert_paragraph_before()
    p_break.add_run().add_break(WD_BREAK.PAGE)
    
    doc.save("test_output.docx")
    print("Saved test_output.docx successfully!")
    if os.path.exists("test_output.docx"):
        os.remove("test_output.docx")

if __name__ == "__main__":
    main()
