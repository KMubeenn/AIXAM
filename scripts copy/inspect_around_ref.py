import docx
import os

def main():
    docs_dir = r"d:\ML\Lawbot\docs"
    r_path = os.path.join(docs_dir, "Constitution and Legal Assistance RAG Bot - FYP Report.docx")
    
    if os.path.exists(r_path):
        doc = docx.Document(r_path)
        print("Paragraphs 170 to 180:")
        for idx in range(170, min(181, len(doc.paragraphs))):
            print(f"[{idx}] {doc.paragraphs[idx].text[:120]}")
    else:
        print("File not found!")

if __name__ == "__main__":
    main()
