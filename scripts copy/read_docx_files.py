import docx
import os

def read_docx(file_path):
    print(f"Reading {file_path}...")
    doc = docx.Document(file_path)
    content = []
    word_count = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Count words
        word_count += len(text.split())
        
        # Check style for headings
        style_name = para.style.name
        if style_name.startswith('Heading'):
            content.append(f"[{style_name}] {text}")
        else:
            # For body text, just sample the first 80 characters
            content.append(f"[Body] {text[:80]}..." if len(text) > 80 else f"[Body] {text}")
            
    # Check tables too
    for i, table in enumerate(doc.tables):
        content.append(f"[Table {i+1}] {len(table.rows)} rows x {len(table.columns)} columns")
        
    return content, word_count

def main():
    docs_dir = r"d:\ML\Lawbot\docs"
    irshad_path = os.path.join(docs_dir, "Irshad_FYP_Mid.docx")
    aixam_path = os.path.join(docs_dir, "FYP Report-AIXAM.docx")
    
    out_file = r"d:\ML\Lawbot\docs\docx_analysis.txt"
    
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("=== ANALYSIS OF DOCX FILES ===\n\n")
        
        if os.path.exists(irshad_path):
            content_irshad, words_irshad = read_docx(irshad_path)
            f.write(f"File: Irshad_FYP_Mid.docx\n")
            f.write(f"Word Count: {words_irshad}\n")
            f.write("Structure:\n")
            for item in content_irshad[:150]:  # Let's show first 150 items
                f.write(item + "\n")
            if len(content_irshad) > 150:
                f.write(f"... and {len(content_irshad) - 150} more paragraphs/headings ...\n")
        else:
            f.write("Irshad_FYP_Mid.docx not found!\n")
            
        f.write("\n" + "="*50 + "\n\n")
        
        if os.path.exists(aixam_path):
            content_aixam, words_aixam = read_docx(aixam_path)
            f.write(f"File: FYP Report-AIXAM.docx\n")
            f.write(f"Word Count: {words_aixam}\n")
            f.write("Structure:\n")
            for item in content_aixam[:300]:  # Show more for AIXAM since it's the complete one
                f.write(item + "\n")
            if len(content_aixam) > 300:
                f.write(f"... and {len(content_aixam) - 300} more paragraphs/headings ...\n")
        else:
            f.write("FYP Report-AIXAM.docx not found!\n")
            
    print(f"Analysis written to {out_file}")

if __name__ == "__main__":
    main()
