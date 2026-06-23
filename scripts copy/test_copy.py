import docx
import os

def copy_paragraph(src_p, dest_p):
    dest_p.alignment = src_p.alignment
    try:
        dest_p.style = src_p.style
    except Exception:
        pass
    dest_p.paragraph_format.space_before = src_p.paragraph_format.space_before
    dest_p.paragraph_format.space_after = src_p.paragraph_format.space_after
    dest_p.paragraph_format.line_spacing = src_p.paragraph_format.line_spacing
    dest_p.paragraph_format.keep_with_next = src_p.paragraph_format.keep_with_next
    
    for run in src_p.runs:
        dest_run = dest_p.add_run(run.text)
        dest_run.bold = run.bold
        dest_run.italic = run.italic
        dest_run.underline = run.underline
        if run.font.name:
            dest_run.font.name = run.font.name
        if run.font.size:
            dest_run.font.size = run.font.size
        try:
            if run.font.color and run.font.color.rgb:
                dest_run.font.color.rgb = run.font.color.rgb
        except Exception:
            pass

def main():
    docs_dir = r"d:\ML\Lawbot\docs"
    src_path = os.path.join(docs_dir, "Constitution and Legal Assistance RAG Bot - FYP Report.docx")
    
    if os.path.exists(src_path):
        src_doc = docx.Document(src_path)
        dest_doc = docx.Document()
        
        # Copy the first 5 paragraphs
        for idx in range(min(5, len(src_doc.paragraphs))):
            src_p = src_doc.paragraphs[idx]
            dest_p = dest_doc.add_paragraph()
            copy_paragraph(src_p, dest_p)
            
        dest_doc.save("test_copy_output.docx")
        print("Copied paragraphs successfully!")
        if os.path.exists("test_copy_output.docx"):
            os.remove("test_copy_output.docx")
    else:
        print("Source file not found!")

if __name__ == "__main__":
    main()
