# Create Chapter 3 SRS DOCX Document for Lawbot
# This script generates a comprehensive DOCX file with Use Cases, SSDs, SRS, and Test Plan

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_chapter3_docx():
    doc = Document()
    
    # Set up styles
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    
    # ============== TITLE PAGE ==============
    title = doc.add_heading('Chapter 3: System Requirements, Architecture & Design', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    subtitle = doc.add_paragraph('Lawbot - Constitution & Legal Assistance RAG System')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph('Final Year Project Documentation')
    doc.add_page_break()
    
    # ============== TABLE OF CONTENTS ==============
    doc.add_heading('Table of Contents', level=1)
    toc_items = [
        '3.1 Project Overview',
        '3.2 Fully Dressed Use Cases',
        '3.3 System Sequence Diagrams (SSD)',
        '3.4 Software Requirements Specification (SRS)',
        '3.5 Test Plan',
    ]
    for item in toc_items:
        doc.add_paragraph(item, style='List Number')
    doc.add_page_break()
    
    # ============== 3.1 PROJECT OVERVIEW ==============
    doc.add_heading('3.1 Project Overview', level=1)
    
    overview_text = """Lawbot is an AI-powered legal assistance chatbot that leverages Retrieval-Augmented Generation (RAG) to provide accurate, context-aware answers to legal and constitutional queries. The system integrates modern NLP technologies with a robust backend infrastructure to deliver seamless legal information retrieval."""
    doc.add_paragraph(overview_text)
    
    doc.add_heading('System Actors', level=2)
    actors = [
        ('User', 'Lawyers, legal professionals, and citizens seeking legal information'),
        ('System', 'The Lawbot application including all backend services'),
        ('LLM Provider', 'External AI services (Groq, Google Gemini) for response generation'),
        ('Database', 'Supabase PostgreSQL for persistent data storage'),
    ]
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Actor'
    hdr_cells[1].text = 'Description'
    for actor, desc in actors:
        row_cells = table.add_row().cells
        row_cells[0].text = actor
        row_cells[1].text = desc
    
    doc.add_paragraph()
    
    # ============== 3.2 USE CASES ==============
    doc.add_heading('3.2 Fully Dressed Use Cases', level=1)
    
    # UC-01: User Registration
    doc.add_heading('UC-01: User Registration', level=2)
    uc01 = [
        ('Use Case ID', 'UC-01'),
        ('Use Case Name', 'User Registration'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User has access to the Lawbot web application'),
        ('Postconditions', 'User account is created and JWT token is issued'),
        ('Main Flow', '1. User navigates to signup page\n2. User enters name, email, and password\n3. System validates input fields\n4. System creates user account in Supabase\n5. System generates JWT authentication token\n6. System redirects user to dashboard'),
        ('Alternative Flows', 'A1: If email already exists, display error message'),
        ('Exception Flows', 'E1: Database connection failure - display error and retry option'),
    ]
    table = doc.add_table(rows=len(uc01), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc01):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # UC-02: User Login
    doc.add_heading('UC-02: User Login', level=2)
    uc02 = [
        ('Use Case ID', 'UC-02'),
        ('Use Case Name', 'User Login'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User has a registered account'),
        ('Postconditions', 'User is authenticated and session is established'),
        ('Main Flow', '1. User enters email and password\n2. System validates credentials against Supabase\n3. System generates JWT token\n4. System loads user profile and chat history\n5. User is redirected to chat interface'),
        ('Alternative Flows', 'A1: Invalid credentials - display error message'),
        ('Exception Flows', 'E1: Authentication service unavailable - display retry option'),
    ]
    table = doc.add_table(rows=len(uc02), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc02):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # UC-03: Text Chat Query
    doc.add_heading('UC-03: Text Chat Query', level=2)
    uc03 = [
        ('Use Case ID', 'UC-03'),
        ('Use Case Name', 'Text Chat Query'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User is authenticated and has an active chat session'),
        ('Postconditions', 'User receives AI-generated legal response with sources'),
        ('Main Flow', '1. User types legal question in chat interface\n2. System performs hybrid RAG search (Vector + BM25 + CrossEncoder)\n3. System retrieves relevant constitutional documents from Pinecone\n4. LLM generates contextual response with retrieved information\n5. Response is streamed token-by-token to user\n6. Message pair is persisted to Supabase'),
        ('Alternative Flows', 'A1: No relevant documents found - LLM provides general guidance\nA2: Primary LLM unavailable - fallback to secondary LLM model'),
        ('Exception Flows', 'E1: RAG pipeline failure - return cached response or error message'),
    ]
    table = doc.add_table(rows=len(uc03), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc03):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # UC-04: Document Upload
    doc.add_heading('UC-04: Document Upload', level=2)
    uc04 = [
        ('Use Case ID', 'UC-04'),
        ('Use Case Name', 'Document Upload'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User is authenticated with active session'),
        ('Postconditions', 'Document content is extracted and available for queries'),
        ('Main Flow', '1. User selects PDF, DOCX, or TXT file\n2. System uploads file to server\n3. DocumentReader extracts text content\n4. Extracted text is added to prompt context\n5. User can query the uploaded document'),
        ('Alternative Flows', 'A1: Unsupported file format - display error with supported formats'),
        ('Exception Flows', 'E1: File too large - display size limit error'),
    ]
    table = doc.add_table(rows=len(uc04), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc04):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # UC-05: Voice Query
    doc.add_heading('UC-05: Voice Query', level=2)
    uc05 = [
        ('Use Case ID', 'UC-05'),
        ('Use Case Name', 'Voice Query'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User has microphone access and is authenticated'),
        ('Postconditions', 'Voice is transcribed and processed as text query'),
        ('Main Flow', '1. User clicks microphone button\n2. System captures audio stream\n3. Audio is sent to VoiceAgent via WebSocket\n4. Whisper model transcribes audio to text\n5. Transcribed text is processed as UC-03 (Text Chat Query)\n6. Response is returned via WebSocket'),
        ('Alternative Flows', 'A1: Low audio quality - request user to repeat'),
        ('Exception Flows', 'E1: Whisper service unavailable - prompt text input instead'),
    ]
    table = doc.add_table(rows=len(uc05), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc05):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # UC-06: Session Management
    doc.add_heading('UC-06: Session Management', level=2)
    uc06 = [
        ('Use Case ID', 'UC-06'),
        ('Use Case Name', 'Session Management'),
        ('Primary Actor', 'User'),
        ('Preconditions', 'User is authenticated'),
        ('Postconditions', 'Chat sessions are created, viewed, or deleted'),
        ('Main Flow', '1. User creates new chat session\n2. System generates unique session ID\n3. User can switch between sessions\n4. System loads session history from Supabase\n5. User can delete old sessions'),
        ('Alternative Flows', 'A1: Maximum sessions reached - prompt to delete old session'),
        ('Exception Flows', 'E1: Session not found - create new session'),
    ]
    table = doc.add_table(rows=len(uc06), cols=2)
    table.style = 'Table Grid'
    for i, (field, value) in enumerate(uc06):
        table.rows[i].cells[0].text = field
        table.rows[i].cells[1].text = value
    doc.add_paragraph()
    
    # ============== 3.3 SYSTEM SEQUENCE DIAGRAMS ==============
    doc.add_heading('3.3 System Sequence Diagrams (SSD)', level=1)
    
    doc.add_heading('SSD-01: Text Chat Query Flow', level=2)
    ssd1_desc = """The following sequence diagram illustrates the text chat query flow:

1. User → Frontend: Submit legal query
2. Frontend → Backend API: POST /chat (session_id, message)
3. Backend → RAG Pipeline: Query vector store
4. RAG Pipeline → Pinecone: Dense vector search (k=20)
5. Pinecone → RAG Pipeline: Return candidate documents
6. RAG Pipeline → RAG Pipeline: BM25 ranking + CrossEncoder reranking
7. RAG Pipeline → Backend: Return top 3 relevant documents
8. Backend → LLM Provider: Generate response with context
9. LLM Provider → Backend: Stream tokens
10. Backend → Frontend: StreamingResponse (token-by-token)
11. Frontend → User: Display response
12. Backend → Supabase: Persist message pair"""
    doc.add_paragraph(ssd1_desc)
    doc.add_paragraph()
    
    doc.add_heading('SSD-02: Voice Query Flow', level=2)
    ssd2_desc = """The voice query sequence diagram:

1. User → Frontend: Click microphone, speak query
2. Frontend → Backend: WebSocket /voice (audio chunks)
3. Backend → VoiceAgent: Accumulate audio bytes
4. VoiceAgent → Whisper: Transcribe audio
5. Whisper → VoiceAgent: Return transcribed text
6. VoiceAgent → ChatBot: Process as text query
7. ChatBot → RAG Pipeline: Retrieve context (same as SSD-01)
8. ChatBot → LLM Provider: Generate response
9. Backend → Frontend: Stream response via WebSocket
10. Frontend → User: Display response"""
    doc.add_paragraph(ssd2_desc)
    doc.add_paragraph()
    
    doc.add_heading('SSD-03: Authentication Flow', level=2)
    ssd3_desc = """The authentication sequence diagram:

1. User → Frontend: Enter credentials (Login/Signup)
2. Frontend → Backend: POST /api/auth/login or /signup
3. Backend → Supabase: Validate/Create user
4. Supabase → Backend: Return user data
5. Backend → Backend: Generate JWT token
6. Backend → Frontend: Return {user, token}
7. Frontend → LocalStorage: Store token
8. Frontend → User: Redirect to dashboard"""
    doc.add_paragraph(ssd3_desc)
    doc.add_paragraph()
    
    # ============== 3.4 SRS ==============
    doc.add_heading('3.4 Software Requirements Specification (SRS)', level=1)
    
    doc.add_heading('3.4.1 Functional Requirements', level=2)
    
    functional_reqs = [
        ('FR-01', 'User Registration', 'System shall allow users to create accounts with name, email, and password'),
        ('FR-02', 'User Authentication', 'System shall authenticate users using JWT tokens with 7-day expiry'),
        ('FR-03', 'Text Chat Interface', 'System shall provide real-time chat interface with streaming responses'),
        ('FR-04', 'RAG-based Retrieval', 'System shall use hybrid search (Vector + BM25 + CrossEncoder) for document retrieval'),
        ('FR-05', 'Document Upload', 'System shall support PDF, DOCX, and TXT file uploads for context'),
        ('FR-06', 'Voice Input', 'System shall transcribe voice queries using Whisper speech-to-text'),
        ('FR-07', 'Session Management', 'System shall support multiple chat sessions per user (max 10)'),
        ('FR-08', 'Message Persistence', 'System shall persist all messages to Supabase PostgreSQL'),
        ('FR-09', 'Chat History', 'System shall display chat history with session titles'),
        ('FR-10', 'Fallback LLM', 'System shall automatically switch to secondary LLM on primary failure'),
        ('FR-11', 'Response Caching', 'System shall cache frequent responses using Redis for optimization'),
        ('FR-12', 'Constitutional Knowledge Base', 'System shall maintain pre-indexed constitution documents in Pinecone'),
        ('FR-13', 'Real-time Cost Tracking', 'System shall track and display LLM API usage costs'),
        ('FR-14', 'Streaming Responses', 'System shall stream LLM responses token-by-token'),
        ('FR-15', 'Profile Management', 'System shall allow users to view and update profile information'),
    ]
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Requirement'
    hdr_cells[2].text = 'Description'
    for req_id, req_name, req_desc in functional_reqs:
        row_cells = table.add_row().cells
        row_cells[0].text = req_id
        row_cells[1].text = req_name
        row_cells[2].text = req_desc
    doc.add_paragraph()
    
    doc.add_heading('3.4.2 Non-Functional Requirements', level=2)
    
    nfr_reqs = [
        ('NFR-01', 'Performance', 'System shall respond to queries within 3 seconds (excluding LLM generation time)'),
        ('NFR-02', 'Availability', 'System shall maintain 99.5% uptime'),
        ('NFR-03', 'Scalability', 'System shall support up to 1000 concurrent users'),
        ('NFR-04', 'Security', 'System shall encrypt all data in transit using HTTPS/TLS'),
        ('NFR-05', 'Usability', 'System shall be accessible on desktop and mobile browsers'),
        ('NFR-06', 'Reliability', 'System shall implement graceful degradation on component failures'),
        ('NFR-07', 'Maintainability', 'System shall follow modular architecture for easy updates'),
        ('NFR-08', 'Data Privacy', 'System shall comply with data protection regulations'),
    ]
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ID'
    hdr_cells[1].text = 'Category'
    hdr_cells[2].text = 'Description'
    for req_id, category, req_desc in nfr_reqs:
        row_cells = table.add_row().cells
        row_cells[0].text = req_id
        row_cells[1].text = category
        row_cells[2].text = req_desc
    doc.add_paragraph()
    
    doc.add_heading('3.4.3 Technology Stack', level=2)
    
    tech_stack = [
        ('Backend Framework', 'Django 5.0+ with FastAPI integration'),
        ('Frontend', 'React 18+ with Vite, TailwindCSS'),
        ('Database', 'Supabase PostgreSQL'),
        ('Vector Database', 'Pinecone (cloud) with FAISS fallback'),
        ('LLM Providers', 'Groq (Llama 3.1), Google Gemini (fallback)'),
        ('Embeddings', 'HuggingFace sentence-transformers/all-MiniLM-L6-v2'),
        ('Speech-to-Text', 'OpenAI Whisper (base model)'),
        ('Caching', 'Redis for response caching'),
        ('Authentication', 'JWT tokens with python-jose'),
    ]
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Component'
    hdr_cells[1].text = 'Technology'
    for component, tech in tech_stack:
        row_cells = table.add_row().cells
        row_cells[0].text = component
        row_cells[1].text = tech
    doc.add_paragraph()
    
    # ============== 3.5 TEST PLAN ==============
    doc.add_heading('3.5 Test Plan', level=1)
    
    doc.add_heading('3.5.1 Test Levels', level=2)
    
    test_levels = [
        ('Unit Testing', 'pytest', 'Test individual functions and classes in isolation (ChatBot, RAGPipeline, VoiceAgent)'),
        ('Integration Testing', 'pytest + httpx', 'Test API endpoints and database interactions'),
        ('System Testing', 'Selenium/Playwright', 'End-to-end testing of complete user workflows'),
        ('User Acceptance Testing', 'Manual', 'Validation with target users (lawyers, legal professionals)'),
    ]
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Test Level'
    hdr_cells[1].text = 'Tools'
    hdr_cells[2].text = 'Description'
    for level, tools, desc in test_levels:
        row_cells = table.add_row().cells
        row_cells[0].text = level
        row_cells[1].text = tools
        row_cells[2].text = desc
    doc.add_paragraph()
    
    doc.add_heading('3.5.2 Testing Techniques', level=2)
    
    techniques = [
        ('Black-box Testing', 'Test system functionality without knowledge of internal implementation'),
        ('White-box Testing', 'Test internal logic and code paths in RAG pipeline and agents'),
        ('Regression Testing', 'Ensure new changes do not break existing functionality'),
        ('Performance Testing', 'Measure response times, throughput, and resource usage'),
        ('Security Testing', 'Validate authentication, authorization, and data protection'),
        ('Load Testing', 'Test system behavior under concurrent user load'),
    ]
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Technique'
    hdr_cells[1].text = 'Description'
    for technique, desc in techniques:
        row_cells = table.add_row().cells
        row_cells[0].text = technique
        row_cells[1].text = desc
    doc.add_paragraph()
    
    doc.add_heading('3.5.3 Test Cases Summary', level=2)
    
    test_cases = [
        ('TC-01', 'FR-01', 'Verify user registration with valid data'),
        ('TC-02', 'FR-01', 'Verify registration fails with duplicate email'),
        ('TC-03', 'FR-02', 'Verify login with valid credentials'),
        ('TC-04', 'FR-02', 'Verify login fails with invalid password'),
        ('TC-05', 'FR-03', 'Verify text chat returns relevant response'),
        ('TC-06', 'FR-04', 'Verify RAG retrieves correct constitutional articles'),
        ('TC-07', 'FR-05', 'Verify PDF document upload and text extraction'),
        ('TC-08', 'FR-06', 'Verify voice transcription accuracy'),
        ('TC-09', 'FR-07', 'Verify session creation and switching'),
        ('TC-10', 'FR-10', 'Verify fallback LLM activation on primary failure'),
        ('TC-11', 'NFR-01', 'Verify response time under 3 seconds'),
        ('TC-12', 'NFR-04', 'Verify HTTPS encryption on all endpoints'),
    ]
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Test Case ID'
    hdr_cells[1].text = 'Requirement'
    hdr_cells[2].text = 'Description'
    for tc_id, req, desc in test_cases:
        row_cells = table.add_row().cells
        row_cells[0].text = tc_id
        row_cells[1].text = req
        row_cells[2].text = desc
    
    # Save document
    doc.save(r'd:\ML\Lawbot\docs\Chapter_3_SRS.docx')
    print("Document created successfully: d:\\ML\\Lawbot\\docs\\Chapter_3_SRS.docx")

if __name__ == "__main__":
    create_chapter3_docx()
