import io
import re
import base64
try:
    import markdown
    MARKDOWN_AVAILABLE=True
except ImportError:
    MARKDOWN_AVAILABLE=False
from datetime import datetime

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt
    DOCX_AVAILABLE=True
except ImportError:
    DOCX_AVAILABLE=False

try:
    from pptx import Presentation
    PPTX_AVAILABLE=True
except ImportError:
    PPTX_AVAILABLE=False

try:
    from fpdf import FPDF, HTMLMixin
    class HTMLFPDF(FPDF, HTMLMixin):
        pass
    PDF_AVAILABLE=True
except ImportError:
    # Fallback for older fpdf2 or different install
    try:
        from fpdf import FPDF
        PDF_AVAILABLE=True
    except ImportError:
        PDF_AVAILABLE=False


class DocumentWriter():

    def _generate_filename(self, title: str, ext: str) -> str:
        safe_title = re.sub(r'[^\w\s-]', '', title).strip().replace(' ', '_')[:50]
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"{safe_title}_{timestamp}.{ext}"

    def write(self, title: str, content: str, format: str) -> dict:
        """Generate a document and return it as a base64-encoded dict."""
        if format == 'pdf':
            return self._write_pdf(title, content)
        elif format == 'docx':
            return self._write_docx(title, content)
        elif format == 'pptx':
            return self._write_pptx(title, content)
        else:
            raise ValueError(f"Unsupported format: {format}")

    def _write_pdf(self, title: str, content: str) -> dict:
        if not PDF_AVAILABLE:
            raise ImportError("fpdf2 is required for PDF generation.")

        # Use HTMLMixin if available for better formatting
        if 'HTMLMixin' in globals() or 'HTMLMixin' in locals() or hasattr(FPDF, 'write_html'):
            pdf = HTMLFPDF() if 'HTMLFPDF' in globals() else FPDF()
        else:
            pdf = FPDF()

        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        
        # Title
        pdf.set_font('Helvetica', 'B', 20)
        pdf.cell(0, 15, title, ln=True, align='C')
        pdf.ln(10)

        # Convert Markdown to HTML for fpdf2's write_html
        # We wrap it in a div with some basic style
        if MARKDOWN_AVAILABLE:
            html_content = markdown.markdown(content)
        else:
            html_content = content.replace('\n', '<br>')
        
        # Basic CSS-like styling for write_html
        full_html = f"""
        <div style="font-family: Helvetica; font-size: 11pt; color: #333333; line-height: 1.5;">
            {html_content}
        </div>
        """
        
        try:
            pdf.write_html(full_html)
        except Exception as e:
            # Fallback to plain text if HTML rendering fails
            pdf.set_font('Helvetica', '', 11)
            for line in content.split('\n'):
                pdf.multi_cell(0, 6, line)

        buffer = io.BytesIO()
        pdf.output(buffer)
        file_bytes = buffer.getvalue()
        buffer.close()

        return {
            'filename': self._generate_filename(title, 'pdf'),
            'file_base64': base64.b64encode(file_bytes).decode('utf-8'),
            'mime_type': 'application/pdf'
        }

    def _write_docx(self, title: str, content: str) -> dict:
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for DOCX generation.")

        doc = DocxDocument()
        doc.add_heading(title, level=0)

        for line in content.split('\n'):
            line = line.strip()
            if not line:
                continue
            
            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('- ') or line.startswith('* '):
                p = doc.add_paragraph(style='List Bullet')
                self._add_formatted_text(p, line[2:])
            else:
                p = doc.add_paragraph()
                self._add_formatted_text(p, line)

        buffer = io.BytesIO()
        doc.save(buffer)
        file_bytes = buffer.getvalue()
        buffer.close()

        return {
            'filename': self._generate_filename(title, 'docx'),
            'file_base64': base64.b64encode(file_bytes).decode('utf-8'),
            'mime_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        }

    def _add_formatted_text(self, paragraph, text):
        """Helper to parse basic markdown (**bold**, *italic*) and add to a docx paragraph."""
        # Split by bold and italic markers
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                paragraph.add_run(part[2:-2]).bold = True
            elif part.startswith('*') and part.endswith('*'):
                paragraph.add_run(part[1:-1]).italic = True
            else:
                paragraph.add_run(part)

    def _write_pptx(self, title: str, content: str) -> dict:
        if not PPTX_AVAILABLE:
            raise ImportError("python-pptx is required for PPTX generation.")

        prs = Presentation()
        
        # Title Slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = title

        # Split content into slides by main headings (# )
        sections = content.split('\n# ')
        for i, section in enumerate(sections):
            if not section.strip(): continue
            
            # Clean up first section if it started with #
            if i == 0 and section.startswith('# '):
                section = section[2:]
            
            lines = section.strip().split('\n')
            slide_title = lines[0].strip().lstrip('# ')
            
            # Use bullet layout
            bullet_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = slide_title
            
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.word_wrap = True
            
            first = True
            for line in lines[1:]:
                clean_line = line.strip().lstrip('-* ').strip()
                if not clean_line: continue
                
                if first:
                    tf.text = clean_line
                    first = False
                else:
                    p = tf.add_paragraph()
                    p.text = clean_line
                    p.level = 0 if not line.startswith('  ') else 1

        buffer = io.BytesIO()
        prs.save(buffer)
        file_bytes = buffer.getvalue()
        buffer.close()

        return {
            'filename': self._generate_filename(title, 'pptx'),
            'file_base64': base64.b64encode(file_bytes).decode('utf-8'),
            'mime_type': 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        }
